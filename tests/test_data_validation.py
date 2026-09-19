"""Per-column data validation.

The dropdown is why anyone reaches for this, so that is what most of these
check. The rest guard the edges Excel itself imposes — a 255-character cap on
an inline list, and integer rules that must not silently swallow a fraction.
"""

import openpyxl
import pytest

from rustpy_xlsxwriter import FastExcel, write_worksheet, write_worksheets

ROWS = [{"status": "open", "qty": 1}, {"status": "closed", "qty": 2}]


def _validations(tmp_path, rules, rows=None, **kwargs):
    path = tmp_path / "dv.xlsx"
    write_worksheet(
        rows if rows is not None else ROWS,
        str(path),
        data_validations=rules,
        **kwargs,
    )
    ws = openpyxl.load_workbook(path).active
    return list(ws.data_validations.dataValidation)


# --- the dropdown -----------------------------------------------------------


def test_list_becomes_a_dropdown(tmp_path):
    (dv,) = _validations(
        tmp_path, {"status": {"type": "list", "values": ["open", "closed"]}}
    )
    assert dv.type == "list"
    assert dv.formula1 == '"open,closed"'


def test_range_covers_the_data_rows_only(tmp_path):
    """Never the header — a header is not a value to validate."""
    (dv,) = _validations(tmp_path, {"status": {"type": "list", "values": ["a"]}})
    assert str(dv.sqref) == "A2:A3"


def test_range_follows_header_row(tmp_path):
    (dv,) = _validations(
        tmp_path, {"status": {"type": "list", "values": ["a"]}}, header_row=2
    )
    assert str(dv.sqref) == "A4:A5"


def test_no_data_rows_writes_nothing(tmp_path):
    assert _validations(tmp_path, {"status": {"type": "list", "values": ["a"]}}, rows=[]) == []


def test_list_over_excels_limit_raises(tmp_path):
    """255 characters including separators; a longer one Excel will not open."""
    with pytest.raises(Exception, match="(?i)length|255|exceed"):
        _validations(tmp_path, {"status": {"type": "list", "values": ["x" * 40] * 10}})


# --- numeric rules ----------------------------------------------------------


def test_whole_number(tmp_path):
    (dv,) = _validations(
        tmp_path, {"qty": {"type": "whole_number", "criteria": ">=", "value": 0}}
    )
    assert dv.type == "whole"
    assert dv.operator == "greaterThanOrEqual"
    assert dv.formula1 == "0"


def test_whole_number_between(tmp_path):
    (dv,) = _validations(
        tmp_path,
        {"qty": {"type": "whole_number", "criteria": "between", "min": 1, "max": 9}},
    )
    # "between" is Excel's default operator, so it is omitted from the XML
    # rather than written out — the two bounds are what prove the rule.
    assert dv.operator in (None, "between")
    assert (dv.formula1, dv.formula2) == ("1", "9")


def test_decimal_keeps_the_fraction(tmp_path):
    (dv,) = _validations(
        tmp_path, {"qty": {"type": "decimal", "criteria": "<", "value": 2.5}}
    )
    assert dv.type == "decimal"
    assert dv.formula1 == "2.5"


def test_whole_number_refuses_a_fraction(tmp_path):
    """Truncating 2.5 to 2 would silently validate against the wrong bound."""
    with pytest.raises(ValueError, match="needs whole numbers; got 2.5"):
        _validations(tmp_path, {"qty": {"type": "whole_number", "criteria": ">", "value": 2.5}})


def test_text_length(tmp_path):
    (dv,) = _validations(
        tmp_path, {"status": {"type": "text_length", "criteria": "<=", "value": 10}}
    )
    assert dv.type == "textLength"


def test_custom_formula(tmp_path):
    (dv,) = _validations(tmp_path, {"qty": {"type": "custom", "formula": "=B2>0"}})
    assert dv.type == "custom"


def test_any_accepts_everything(tmp_path):
    (dv,) = _validations(
        tmp_path, {"qty": {"type": "any", "input_message": "anything goes"}}
    )
    assert dv.prompt == "anything goes"


# --- messages ---------------------------------------------------------------


def test_input_and_error_messages(tmp_path):
    (dv,) = _validations(
        tmp_path,
        {"status": {
            "type": "list", "values": ["open"],
            "input_title": "Status", "input_message": "Pick one",
            "error_title": "Nope", "error_message": "Not a status",
        }},
    )
    assert (dv.promptTitle, dv.prompt) == ("Status", "Pick one")
    assert (dv.errorTitle, dv.error) == ("Nope", "Not a status")


def test_error_style(tmp_path):
    (dv,) = _validations(
        tmp_path,
        {"status": {"type": "list", "values": ["open"], "error_style": "warning"}},
    )
    assert dv.errorStyle == "warning"


def test_unknown_error_style_raises(tmp_path):
    with pytest.raises(ValueError, match="unknown error_style 'shout'"):
        _validations(
            tmp_path,
            {"status": {"type": "list", "values": ["open"], "error_style": "shout"}},
        )


# --- validation of the spec itself ------------------------------------------


def test_unknown_column_warns_and_is_skipped(tmp_path):
    with pytest.warns(UserWarning, match="data_validations: unknown column 'nope'"):
        assert _validations(tmp_path, {"nope": {"type": "list", "values": ["a"]}}) == []


def test_unknown_type_raises(tmp_path):
    with pytest.raises(ValueError, match="unknown type 'dropdown'"):
        _validations(tmp_path, {"status": {"type": "dropdown", "values": ["a"]}})


def test_unknown_criteria_raises(tmp_path):
    with pytest.raises(ValueError, match="unknown criteria '=>'"):
        _validations(tmp_path, {"qty": {"type": "decimal", "criteria": "=>", "value": 1}})


def test_list_needs_values(tmp_path):
    with pytest.raises(ValueError, match="a 'list' rule needs 'values'"):
        _validations(tmp_path, {"status": {"type": "list"}})


def test_between_needs_max(tmp_path):
    with pytest.raises(ValueError, match="a 'decimal' rule needs 'max'"):
        _validations(
            tmp_path, {"qty": {"type": "decimal", "criteria": "between", "min": 1}}
        )


def test_must_be_a_dict(tmp_path):
    with pytest.raises(ValueError, match="must be a dict keyed by column name"):
        _validations(tmp_path, ["status"])


def test_rule_must_be_a_dict(tmp_path):
    with pytest.raises(ValueError, match="rule for 'status' must be a dict"):
        _validations(tmp_path, {"status": ["open", "closed"]})


# --- every entry point ------------------------------------------------------


def test_multi_sheet_is_keyed_by_sheet_name(tmp_path):
    path = tmp_path / "multi.xlsx"
    write_worksheets(
        [("A", ROWS), ("B", ROWS)],
        str(path),
        data_validations={"A": {"status": {"type": "list", "values": ["open"]}}},
    )
    book = openpyxl.load_workbook(path)
    assert len(list(book["A"].data_validations.dataValidation)) == 1
    assert list(book["B"].data_validations.dataValidation) == []


def test_builder(tmp_path):
    path = tmp_path / "b.xlsx"
    FastExcel(path).sheet(
        "S", ROWS, data_validations={"status": {"type": "list", "values": ["open"]}}
    ).save()
    ws = openpyxl.load_workbook(path).active
    assert len(list(ws.data_validations.dataValidation)) == 1
