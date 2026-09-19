"""Concurrent writes.

`save_workbook` releases the GIL around the XML assembly and the deflate,
which is about two thirds of a write and touches no Python object. That is
what lets two threads overlap — and it is also a new way to be wrong, since
nothing before this ran two writes at once.

These check the part that would break if a buffer, a workbook or a format
were shared across that boundary: every worker's file must contain exactly
its own rows, whole and in order.
"""

import io
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor

import openpyxl
import pytest

from rustpy_xlsxwriter import FastExcel, Format, write_csv, write_worksheet

WORKERS = 8
ROWS_EACH = 400


def _rows(worker):
    """Rows whose values identify the worker that wrote them."""
    return [
        {"worker": worker, "n": i, "label": f"w{worker}-r{i}"}
        for i in range(ROWS_EACH)
    ]


def _read(path):
    ws = openpyxl.load_workbook(path).active
    return [
        (ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value)
        for r in range(2, ws.max_row + 1)
    ]


def test_parallel_writes_do_not_mix(tmp_path):
    """Each file holds its own rows and no other worker's."""
    def job(worker):
        path = tmp_path / f"w{worker}.xlsx"
        write_worksheet(_rows(worker), str(path))
        return worker, path

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        done = list(pool.map(job, range(WORKERS)))

    for worker, path in done:
        got = _read(path)
        assert len(got) == ROWS_EACH
        assert got == [(worker, i, f"w{worker}-r{i}") for i in range(ROWS_EACH)]


def test_parallel_writes_to_buffers(tmp_path):
    """The buffer path hands bytes back to Python after the detached save."""
    def job(worker):
        buf = io.BytesIO()
        write_worksheet(_rows(worker), buf)
        return worker, buf.getvalue()

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        done = list(pool.map(job, range(WORKERS)))

    for worker, data in done:
        assert zipfile.is_zipfile(io.BytesIO(data))
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        assert ws.cell(2, 1).value == worker
        assert ws.cell(ROWS_EACH + 1, 3).value == f"w{worker}-r{ROWS_EACH - 1}"


def test_parallel_csv_writes(tmp_path):
    """CSV takes the same target-writing path at the end."""
    def job(worker):
        buf = io.BytesIO()
        write_csv(_rows(worker), buf)
        return worker, buf.getvalue().decode()

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        done = list(pool.map(job, range(WORKERS)))

    for worker, text in done:
        lines = text.splitlines()
        assert lines[0] == "worker,n,label"
        assert lines[1] == f"{worker},0,w{worker}-r0"
        assert len(lines) == ROWS_EACH + 1


def test_a_shared_format_is_safe_across_threads(tmp_path):
    """One Format applied by many writers at once."""
    shared = Format().set_bold()

    def job(worker):
        path = tmp_path / f"fmt{worker}.xlsx"
        write_worksheet(_rows(worker), str(path), column_formats={"label": shared})
        return path

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        paths = list(pool.map(job, range(WORKERS)))

    for path in paths:
        assert openpyxl.load_workbook(path).active.cell(2, 3).font.bold


def test_parallel_builder_writes(tmp_path):
    """The same through FastExcel, which carries more per-sheet state."""
    def job(worker):
        path = tmp_path / f"b{worker}.xlsx"
        (
            FastExcel(path)
            .format(bold_headers=True)
            .sheet("S", _rows(worker), autofilter=True)
            .save()
        )
        return worker, path

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        done = list(pool.map(job, range(WORKERS)))

    for worker, path in done:
        ws = openpyxl.load_workbook(path)["S"]
        assert ws.cell(2, 1).value == worker
        assert ws.max_row == ROWS_EACH + 1


def test_an_exception_in_one_thread_does_not_disturb_the_others(tmp_path):
    """A failing write must not leave the GIL or another writer in a bad state."""
    def job(worker):
        if worker == 3:
            with pytest.raises(ValueError):
                write_worksheet(
                    _rows(worker), str(tmp_path / "bad.xlsx"), sheet_name="a" * 40
                )
            return None
        path = tmp_path / f"ok{worker}.xlsx"
        write_worksheet(_rows(worker), str(path))
        return worker, path

    with ThreadPoolExecutor(max_workers=WORKERS) as pool:
        done = [d for d in pool.map(job, range(WORKERS)) if d]

    assert len(done) == WORKERS - 1
    for worker, path in done:
        assert _read(path) == [
            (worker, i, f"w{worker}-r{i}") for i in range(ROWS_EACH)
        ]


def test_many_rounds(tmp_path):
    """Repeated, to give a race more than one chance to show itself."""
    errors = []

    def job(worker):
        try:
            for round_ in range(5):
                buf = io.BytesIO()
                write_worksheet(_rows(worker), buf)
                ws = openpyxl.load_workbook(io.BytesIO(buf.getvalue())).active
                assert ws.cell(2, 1).value == worker, (worker, round_)
        except Exception as exc:  # noqa: BLE001 - reported below
            errors.append(exc)

    threads = [threading.Thread(target=job, args=(w,)) for w in range(WORKERS)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()

    assert not errors, errors
