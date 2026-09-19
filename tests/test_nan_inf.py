"""What happens to missing values and the infinities.

Excel has no cell type for either, so they have always been written as an empty
cell — silently, and documented nowhere, which means a column of NaN arrived as
a column of blanks indistinguishable from missing data. That default is kept
(changing it would rewrite every existing file), but it is now stated and can
be overridden, on every write path.
"""

import io

import openpyxl
import pytest

from rustpy_xlsxwriter import FastExcel, write_csv, write_worksheet

NAN, INF, NEG_INF = float("nan"), float("inf"), float("-inf")
ROWS = [{"nan": NAN, "inf": INF, "neg": NEG_INF, "ok": 1.5}]


def _csv(**kwargs):
    buf = io.BytesIO()
    write_csv(ROWS, buf, **kwargs)
    return buf.getvalue().decode()


def _cells(tmp_path, **kwargs):
    path = tmp_path / "out.xlsx"
    write_worksheet(ROWS, str(path), **kwargs)
    ws = openpyxl.load_workbook(path).active
    return [ws.cell(2, c).value for c in range(1, 5)]


# --- the default, now pinned ------------------------------------------------


def test_csv_default_is_empty_fields():
    assert _csv() == "nan,inf,neg,ok\n,,,1.5\n"


def test_excel_default_is_empty_cells(tmp_path):
    assert _cells(tmp_path) == [None, None, None, 1.5]


# --- overrides --------------------------------------------------------------


def test_csv_overrides():
    assert _csv(na_rep="NA", inf_value="INF") == "nan,inf,neg,ok\nNA,INF,-INF,1.5\n"


def test_excel_overrides(tmp_path):
    assert _cells(tmp_path, na_rep="NA", inf_value="INF") == [
        "NA",
        "INF",
        "-INF",
        1.5,
    ]


def test_nan_and_inf_are_independent():
    """Setting one must not conjure a representation for the other."""
    assert _csv(na_rep="NA") == "nan,inf,neg,ok\nNA,,,1.5\n"
    assert _csv(inf_value="INF") == "nan,inf,neg,ok\n,INF,-INF,1.5\n"


def test_negative_infinity_takes_the_sign_from_the_value():
    assert _csv(inf_value="∞") == "nan,inf,neg,ok\n,∞,-∞,1.5\n"


def test_representation_is_escaped_like_any_other_field():
    """It reaches the same escaping as real data, not raw into the buffer."""
    assert _csv(na_rep="a,b") == 'nan,inf,neg,ok\n"a,b",,,1.5\n'


# --- every input path -------------------------------------------------------


def test_arrow_csv_path():
    pd = pytest.importorskip("pandas")
    pytest.importorskip("pyarrow")
    buf = io.BytesIO()
    write_csv(pd.DataFrame(ROWS), buf, na_rep="NA", inf_value="INF")
    assert buf.getvalue().decode() == "nan,inf,neg,ok\nNA,INF,-INF,1.5\n"


def test_arrow_excel_path(tmp_path):
    pd = pytest.importorskip("pandas")
    pytest.importorskip("pyarrow")
    path = tmp_path / "df.xlsx"
    write_worksheet(pd.DataFrame(ROWS), str(path), na_rep="NA", inf_value="INF")
    ws = openpyxl.load_workbook(path).active
    assert [ws.cell(2, c).value for c in range(1, 5)] == ["NA", "INF", "-INF", 1.5]


def test_multi_sheet_applies_to_every_sheet(tmp_path):
    from rustpy_xlsxwriter import write_worksheets

    path = tmp_path / "multi.xlsx"
    write_worksheets(
        [("A", ROWS), ("B", ROWS)], str(path), na_rep="NA", inf_value="INF"
    )
    book = openpyxl.load_workbook(path)
    for name in ("A", "B"):
        assert book[name].cell(2, 1).value == "NA"
        assert book[name].cell(2, 3).value == "-INF"


# --- builder ----------------------------------------------------------------


def test_builder_csv():
    buf = io.BytesIO()
    (
        FastExcel(buf, output_format="csv")
        .format(na_rep="NA", inf_value="INF")
        .sheet("S", ROWS)
        .save()
    )
    assert buf.getvalue().decode() == "nan,inf,neg,ok\nNA,INF,-INF,1.5\n"


def test_builder_excel(tmp_path):
    path = tmp_path / "b.xlsx"
    FastExcel(path).format(na_rep="NA").sheet("S", ROWS).save()
    ws = openpyxl.load_workbook(path).active
    assert ws.cell(2, 1).value == "NA"


def test_csv_does_not_warn_about_them():
    """They are honoured on CSV, so they are not Excel-only options."""
    import warnings

    buf = io.BytesIO()
    with warnings.catch_warnings(record=True) as log:
        warnings.simplefilter("always")
        FastExcel(buf, output_format="csv").format(na_rep="NA").sheet(
            "S", ROWS
        ).save()
    assert not [w for w in log if "Excel-only" in str(w.message)]


def test_na_rep_covers_none_and_nan_alike():
    """The unification, pinned: one knob, three shapes of missing.

    A records ``None``, a float ``NaN`` and an Arrow null must all render the
    same — otherwise the setting works or not depending on which input type the
    caller happened to pass, which is the trap it exists to avoid.
    """
    buf = io.BytesIO()
    write_csv([{"a": None, "b": NAN}], buf, na_rep="NA")
    assert buf.getvalue().decode() == "a,b\nNA,NA\n"


def test_na_rep_covers_arrow_nulls(tmp_path):
    pl = pytest.importorskip("polars")
    buf = io.BytesIO()
    write_csv(pl.DataFrame({"a": [None, 1]}), buf, na_rep="NA")
    assert buf.getvalue().decode() == "a\nNA\n1\n"
